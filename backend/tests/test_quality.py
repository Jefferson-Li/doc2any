"""Explainable conversion quality score."""

from __future__ import annotations

import io
import zipfile

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image

from app.converters import quality
from app.converters.engine import pdf_to_docx_visual
from app.converters.quality import (
    _aspect_ratio_score,
    _clamp,
    _docx_media_images,
    _docx_stats,
    _docx_text,
    _element_alignment,
    _image_preservation,
    _image_similarity,
    _layout_similarity,
    _normalize,
    _page_structure,
    _pdf_image_count,
    _pdf_text,
    _render_pdf_pages,
    _spaced_out_ratio,
    _text_overlap,
    _text_preservation,
    pick_best_candidate,
    score_conversion,
)
from tests.support import add_page, solid_image, write_pdf


def _quality(overall, layout, text, image=100, align=80, page=90):
    return {
        "overall": overall,
        "grade": "B",
        "layout_similarity": layout,
        "text_preservation": text,
        "image_preservation": image,
        "element_alignment": align,
        "page_structure": page,
    }


def _candidate(mode, overall, layout, text):
    return {"mode": mode, "path": f"/{mode}.docx", "quality": _quality(overall, layout, text)}


def _pdf(path, pages=("Hello quality world",)):
    def populate(doc):
        for text in pages:
            if callable(text):
                text(doc)
            else:
                add_page(doc, text)

    return write_pdf(path, populate)


def _docx(path, paragraphs=("Hello quality world",), tables=()):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if tables:
        table = document.add_table(rows=len(tables), cols=len(tables[0]))
        for r, row in enumerate(tables):
            for c, value in enumerate(row):
                table.rows[r].cells[c].text = value
    document.save(path)
    return path


def test_text_helpers(tmp_path):
    assert _clamp(-5) == 0
    assert _clamp(140) == 100
    assert _clamp(12.5) == 12.5
    assert _normalize("  Hello\nWorld ") == "hello world"
    assert _spaced_out_ratio("   ") == 0
    assert _spaced_out_ratio("Hello") == 0
    assert _spaced_out_ratio("H e l l o") > 0
    assert _spaced_out_ratio("你 好 世 界") > 0
    assert _text_overlap("", "") == 1
    assert _text_overlap("a", "") == 0
    assert _text_overlap("", "a") == 0
    assert _text_overlap("a b", "b c") == pytest.approx(1 / 3)

    pdf = _pdf(tmp_path / "two.pdf", pages=("Alpha page", "Beta page"))
    assert "Alpha" in _pdf_text(pdf)
    assert "Beta" in _pdf_text(pdf)
    assert _pdf_image_count(pdf) == 0

    image = solid_image(tmp_path / "pic.png")

    def with_image(doc):
        page = add_page(doc, "caption")
        page.insert_image(fitz.Rect(100, 100, 180, 180), filename=str(image))

    pictured = write_pdf(tmp_path / "img.pdf", with_image)
    assert _pdf_image_count(pictured) >= 1

    docx = _docx(tmp_path / "t.docx", paragraphs=("Para", ""), tables=(("A", "B"),))
    document = Document(str(docx))
    run = document.paragraphs[0].add_run("BoxText")._r
    empty = OxmlElement("w:t")
    empty.text = ""
    run.append(empty)
    document.save(docx)
    extracted = _docx_text(docx)
    assert "Para" in extracted
    assert "A" in extracted and "B" in extracted
    assert "BoxText" in extracted
    stats = _docx_stats(docx)
    assert stats["paragraphs"] >= 1
    assert stats["tables"] == 1
    assert stats["sections"] >= 1


def test_media_images_and_similarity(tmp_path):
    docx = _docx(tmp_path / "media.docx")
    png = io.BytesIO()
    Image.new("RGB", (20, 20), (255, 0, 0)).save(png, format="PNG")
    jpg = io.BytesIO()
    Image.new("RGB", (20, 16), (0, 0, 255)).save(jpg, format="JPEG")
    with zipfile.ZipFile(docx, "a") as archive:
        archive.writestr("word/media/ok.png", png.getvalue())
        archive.writestr("word/media/ok.jpg", jpg.getvalue())
        archive.writestr("word/media/bad.png", b"not-an-image")
        archive.writestr("word/media/readme.txt", b"skip")
        archive.writestr("word/embeddings/other.png", png.getvalue())

    images = _docx_media_images(docx)
    assert len(images) == 2
    assert images[0].mode == "RGB"

    red = Image.new("RGB", (32, 32), (255, 0, 0))
    black = Image.new("RGB", (32, 32), (0, 0, 0))
    white = Image.new("RGB", (32, 32), (255, 255, 255))
    assert _image_similarity(red, red) == pytest.approx(1)
    assert _image_similarity(white, black) == pytest.approx(0)
    assert _image_similarity(red, black) == pytest.approx(1 - (255 / 3) / 255)
    assert _image_similarity(Image.new("RGB", (4, 20), "white"), Image.new("RGB", (4, 20), "white")) == 0
    assert _image_similarity(Image.new("RGB", (20, 4), "white"), Image.new("RGB", (20, 4), "white")) == 0
    assert _aspect_ratio_score(0, 10, 10, 10) == 0
    assert _aspect_ratio_score(100, 100, 100, 100) == pytest.approx(1)
    assert _aspect_ratio_score(100, 100, 300, 100) == 0

    many = write_pdf(
        tmp_path / "nine.pdf",
        lambda doc: [add_page(doc, f"page {i}") for i in range(9)],
    )
    assert len(_render_pdf_pages(many)) == 8
    assert len(_render_pdf_pages(many, max_pages=2)) == 2
    assert len(_render_pdf_pages(_pdf(tmp_path / "one.pdf"), max_pages=3)) == 1


def test_layout_similarity_branches(tmp_path, monkeypatch):
    pdf = _pdf(tmp_path / "s.pdf")
    visual = tmp_path / "visual.docx"
    pdf_to_docx_visual(pdf, visual)
    stats = _docx_stats(visual)
    score, detail = _layout_similarity(
        pdf,
        visual,
        mode_used="visual",
        layout={"layout_type": "simple", "pages": 1},
        spaced=0,
        overlap=0.2,
        stats=stats,
    )
    assert detail["method"] == "visual"
    assert detail["pdf_pages_compared"] == 1
    assert score >= 88

    positioned, positioned_detail = _layout_similarity(
        pdf,
        visual,
        mode_used="positioned",
        layout={"layout_type": "simple", "pages": 1},
        spaced=0,
        overlap=1,
        stats=stats,
    )
    low_overlap, _ = _layout_similarity(
        pdf,
        visual,
        mode_used="positioned",
        layout={"layout_type": "simple", "pages": 1},
        spaced=0,
        overlap=0,
        stats=stats,
    )
    assert positioned_detail["method"] == "positioned"
    assert positioned >= low_overlap

    plain = _docx(tmp_path / "plain.docx", paragraphs=("unrelated",))
    empty_media, empty_detail = _layout_similarity(
        pdf,
        plain,
        mode_used="visual",
        layout={},
        spaced=0,
        overlap=0,
        stats=_docx_stats(plain),
    )
    assert empty_media == 90
    assert empty_detail["page_image_similarity"] == 0.85
    assert empty_detail["aspect_ratio_score"] == 0.9

    two = _pdf(tmp_path / "two.pdf", pages=("one", "two"))
    partial, partial_detail = _layout_similarity(
        two,
        visual,
        mode_used="visual",
        layout={"pages": 2},
        spaced=0,
        overlap=0,
        stats=stats,
    )
    assert partial_detail["pdf_pages_compared"] == 2
    assert partial_detail["docx_media_images"] == 1
    assert partial >= 88

    monkeypatch.setattr(quality, "_render_pdf_pages", lambda *args, **kwargs: [])
    missing, missing_detail = _layout_similarity(
        pdf,
        visual,
        mode_used="visual",
        layout={},
        spaced=0,
        overlap=0,
        stats=stats,
    )
    assert missing == 90
    assert missing_detail["note"] == "empty_pdf"

    heuristic, heuristic_detail = _layout_similarity(
        pdf,
        plain,
        mode_used="flow",
        layout={"layout_type": "simple", "pages": 1},
        spaced=0,
        overlap=1,
        stats={"sections": 1, "paragraphs": 1, "tables": 0, "inline_shapes": 0},
    )
    assert heuristic_detail["method"] == "heuristic_editable"
    assert heuristic == pytest.approx(90)

    penalized, penalized_detail = _layout_similarity(
        pdf,
        plain,
        mode_used="editable",
        layout={"layout_type": "designed", "pages": 1},
        spaced=0,
        overlap=1,
        stats={"sections": 10, "paragraphs": 1, "tables": 0, "inline_shapes": 0},
    )
    assert penalized_detail["complex_layout_penalty"] is True
    assert penalized_detail["section_explosion"] == 10
    assert penalized == pytest.approx(30)

    collapsed, _ = _layout_similarity(
        pdf,
        plain,
        mode_used="editable",
        layout={"layout_type": "simple", "pages": 1},
        spaced=1,
        overlap=0,
        stats={"sections": 1, "paragraphs": 1, "tables": 0, "inline_shapes": 0},
    )
    assert collapsed == 0


def test_metric_components(tmp_path):
    assert _text_preservation(
        mode_used="visual", pdf_chars=10, docx_chars=0, overlap=0, spaced=0, is_scanned=False
    )[0] == 96
    scanned_score, scanned_detail = _text_preservation(
        mode_used="visual", pdf_chars=0, docx_chars=0, overlap=0, spaced=0, is_scanned=True
    )
    assert scanned_score == 96
    assert scanned_detail["method"] == "visual_raster_preservation"
    assert _text_preservation(
        mode_used="visual", pdf_chars=0, docx_chars=0, overlap=0, spaced=0, is_scanned=False
    )[0] == 80

    positioned, positioned_detail = _text_preservation(
        mode_used="positioned",
        pdf_chars=10,
        docx_chars=10,
        overlap=1,
        spaced=0,
        is_scanned=False,
    )
    assert positioned == 100
    assert positioned_detail["method"] == "positioned_textbox_overlap"
    flowing, flowing_detail = _text_preservation(
        mode_used="flow",
        pdf_chars=0,
        docx_chars=0,
        overlap=0,
        spaced=1,
        is_scanned=False,
    )
    assert flowing == 0
    assert flowing_detail["method"] == "token_overlap_and_coverage"
    assert flowing_detail["char_coverage"] == 1

    pdf = _pdf(tmp_path / "s.pdf")
    image = solid_image(tmp_path / "p.png")

    def draw(doc):
        page = add_page(doc, "page one")
        page.insert_image(fitz.Rect(72, 120, 140, 180), filename=str(image))
        add_page(doc, "page two")

    pictured = write_pdf(tmp_path / "pic.pdf", draw)
    visual = tmp_path / "v.docx"
    pdf_to_docx_visual(pdf, visual)
    full, _ = _image_preservation(pdf, visual, mode_used="visual")
    assert full == 100
    partial, partial_detail = _image_preservation(pictured, visual, mode_used="positioned")
    assert partial_detail["pdf_pages"] == 2
    assert 0 < partial < 100
    bare = _docx(tmp_path / "bare.docx", paragraphs=("x",))
    assert _image_preservation(pdf, bare, mode_used="visual")[0] == 40
    assert _image_preservation(pdf, bare, mode_used="flow")[1]["note"] == "source_has_no_images"
    ratio, _ = _image_preservation(pictured, visual, mode_used="flow")
    assert 0 < ratio <= 100

    assert _element_alignment(
        mode_used="visual",
        spaced=0,
        stats={"sections": 1},
        layout={},
        layout_sim_detail={"aspect_ratio_score": 1},
    )[0] == pytest.approx(100)
    assert _element_alignment(
        mode_used="visual",
        spaced=0,
        stats={"sections": 1},
        layout={},
        layout_sim_detail={},
    )[0] == pytest.approx(95.5)
    positioned_align, positioned_align_detail = _element_alignment(
        mode_used="positioned",
        spaced=0,
        stats={"sections": 1},
        layout={},
        layout_sim_detail={"aspect_ratio_score": 0.2},
    )
    # 0 is falsy, so a zero aspect falls back to 0.9 before the positioned bonus.
    falsy_aspect, _ = _element_alignment(
        mode_used="positioned",
        spaced=0,
        stats={"sections": 1},
        layout={},
        layout_sim_detail={"aspect_ratio_score": 0},
    )
    assert positioned_align == pytest.approx(67)
    assert falsy_aspect == pytest.approx(98.5)
    assert positioned_align_detail["method"] == "positioned_textbox_alignment"
    editable_align, editable_detail = _element_alignment(
        mode_used="flow",
        spaced=1,
        stats={"sections": 8},
        layout={"pages": 1, "columns_estimate": 2},
        layout_sim_detail={},
    )
    assert editable_align == pytest.approx(5)
    assert editable_detail["method"] == "editable_flow_alignment"
    plain_align, _ = _element_alignment(
        mode_used="flow",
        spaced=0,
        stats={"sections": 1},
        layout={"pages": 1, "columns_estimate": 1},
        layout_sim_detail={},
    )
    assert plain_align == pytest.approx(90)

    assert _page_structure(pdf, mode_used="visual", stats={"sections": 1, "paragraphs": 1, "tables": 0})[0] == 98
    assert _page_structure(pictured, mode_used="positioned", stats={"sections": 1, "paragraphs": 1, "tables": 0})[0] == 94
    assert _page_structure(pdf, mode_used="visual", stats={"sections": 0, "paragraphs": 0, "tables": 0})[0] == 70
    assert _page_structure(pdf, mode_used="flow", stats={"sections": 1, "paragraphs": 0, "tables": 0})[0] == 20
    exploded, exploded_detail = _page_structure(
        pdf, mode_used="flow", stats={"sections": 5, "paragraphs": 2, "tables": 0}
    )
    assert exploded == 55
    assert exploded_detail["note"] == "section_explosion"
    assert _page_structure(pdf, mode_used="flow", stats={"sections": 1, "paragraphs": 2, "tables": 1})[0] == 92
    assert _page_structure(pdf, mode_used="flow", stats={"sections": 0, "paragraphs": 2, "tables": 0})[0] == 70


def test_score_conversion_grades_weights_and_editability(tmp_path, monkeypatch):
    pdf = _pdf(tmp_path / "s.pdf", pages=("Hello quality world " * 5,))
    docx = _docx(tmp_path / "s.docx")
    visual = tmp_path / "v.docx"
    pdf_to_docx_visual(pdf, visual)

    real = score_conversion(
        pdf,
        docx,
        mode_used="flow",
        layout={"layout_type": "simple", "pages": 1, "is_scanned": False, "columns_estimate": 1},
    )
    assert set(real["explain"]) == {
        "layout_similarity",
        "text_preservation",
        "image_preservation",
        "element_alignment",
        "page_structure",
    }
    assert real["weights"]["text_preservation"] == 0.3
    assert real["thresholds"]["layout_retry"] == 70
    assert real["grade"] in {"A", "B", "C", "D"}
    assert 0 <= real["overall"] <= 100

    scanned = score_conversion(
        pdf,
        visual,
        mode_used="visual",
        layout={"layout_type": "scanned", "pages": 1, "is_scanned": True},
    )
    assert scanned["weights"]["layout_similarity"] == 0.35
    assert scanned["editability"] == 22
    assert scanned["text_preservation"] == 96

    designed = score_conversion(
        pdf,
        visual,
        mode_used="positioned",
        layout={"layout_type": "complex", "pages": 1, "is_scanned": False},
    )
    assert designed["editability"] == 88
    assert designed["weights"]["layout_similarity"] == 0.35

    state = {"value": 90.0}

    def fixed(*args, **kwargs):
        return state["value"], {"aspect_ratio_score": 1}

    for name in (
        "_layout_similarity",
        "_text_preservation",
        "_image_preservation",
        "_element_alignment",
        "_page_structure",
    ):
        monkeypatch.setattr(quality, name, fixed)

    def grade_for(value, mode, layout):
        state["value"] = value
        return score_conversion(pdf, docx, mode_used=mode, layout=layout)["grade"]

    simple = {"layout_type": "simple", "pages": 1}
    assert grade_for(85, "flow", simple) == "A"
    assert grade_for(84.9, "flow", simple) == "B"
    assert grade_for(70, "flow", simple) == "B"
    assert grade_for(69.9, "flow", simple) == "C"
    assert grade_for(55, "flow", simple) == "C"
    assert grade_for(54.9, "flow", None) == "D"

    monkeypatch.setattr(quality, "_spaced_out_ratio", lambda text: 0.5)
    edited = score_conversion(pdf, docx, mode_used="flow", layout=simple)
    assert edited["editability"] == 67.0
    monkeypatch.setattr(quality, "_spaced_out_ratio", lambda text: 2)
    assert score_conversion(pdf, docx, mode_used="editable", layout=simple)["editability"] == 0


def test_pick_best_candidate_rules():
    with pytest.raises(ValueError, match="No conversion candidates"):
        pick_best_candidate([], {"layout_type": "simple"})

    visual = _candidate("visual", 80, 90, 90)
    editable = _candidate("editable", 80, 80, 80)
    scanned = pick_best_candidate([editable, visual], {"layout_type": "scanned"})
    assert scanned["mode"] == "visual"
    assert scanned["retried_from"] == "editable"
    assert scanned["retry_reason"] == "scanned_pdf_no_text_layer"

    only_visual = pick_best_candidate([visual], {"layout_type": "scanned"})
    assert only_visual["retried_from"] is None

    good = pick_best_candidate(
        [_candidate("other", 10, 10, 10), _candidate("positioned", 70, 70, 70), visual],
        {},
    )
    assert good["mode"] == "positioned"
    assert "retried_from" not in good

    ranked_editable = pick_best_candidate(
        [_candidate("editable", 80, 80, 39), _candidate("visual", 90, 90, 90)],
        {"layout_type": "simple"},
    )
    assert ranked_editable["mode"] == "editable"
    assert "retried_from" not in ranked_editable

    low_text = pick_best_candidate(
        [_candidate("positioned", 80, 80, 10), visual],
        {"layout_type": "simple"},
    )
    assert low_text["mode"] == "positioned"

    low_layout = pick_best_candidate(
        [_candidate("editable", 80, 20, 10), _candidate("visual", 10, 10, 10)],
        {"layout_type": "simple"},
    )
    assert low_layout["mode"] == "editable"

    visual_not_better = pick_best_candidate(
        [_candidate("positioned", 10, 20, 10), _candidate("visual", 5, 20, 90)],
        {"layout_type": "simple"},
    )
    assert visual_not_better["mode"] == "positioned"

    only_editable = pick_best_candidate(
        [_candidate("positioned", 10, 10, 10)],
        {"layout_type": "simple"},
    )
    assert only_editable["mode"] == "positioned"

    # text is high enough to miss the first check, overall is not.
    assert pick_best_candidate(
        [_candidate("editable", 40, 80, 50)],
        {"layout_type": "simple"},
    )["mode"] == "editable"
    # overall clears the check, layout similarity does not.
    assert pick_best_candidate(
        [_candidate("positioned", 90, 20, 50)],
        {"layout_type": "simple"},
    )["mode"] == "positioned"

    fallback = pick_best_candidate(
        [
            _candidate("positioned", 0, 10, 10),
            _candidate("visual", 20, 80, 90),
        ],
        {"layout_type": "complex"},
    )
    assert fallback["mode"] == "visual"
    assert fallback["retried_from"] == "positioned"
    assert "fallback" in fallback["retry_reason"]

    visual_wins = pick_best_candidate(
        [_candidate("editable", 0, 0, 0), _candidate("visual", 100, 100, 100)],
        {"layout_type": "simple"},
    )
    assert visual_wins["mode"] == "visual"
    assert "retried_from" not in visual_wins

    only = pick_best_candidate([_candidate("visual", 40, 40, 40)], {"layout_type": "simple"})
    assert only["mode"] == "visual"
    assert "retried_from" not in only

    other = pick_best_candidate(
        [_candidate("other", 15, 15, 15)],
        {"layout_type": "simple"},
    )
    assert other["mode"] == "other"

    mixed_rank = pick_best_candidate(
        [_candidate("other", 100, 10, 10), _candidate("visual", 10, 10, 10)],
        {"layout_type": "simple"},
    )
    assert mixed_rank["mode"] == "visual"

    no_visual = pick_best_candidate(
        [_candidate("editable", 0, 0, 0)],
        {"layout_type": "scanned"},
    )
    assert no_visual["mode"] == "editable"
