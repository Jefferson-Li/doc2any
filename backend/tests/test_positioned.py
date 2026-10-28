"""Positioned editable PDF→DOCX helpers and conversion."""

from __future__ import annotations

import zipfile
from io import BytesIO
from types import SimpleNamespace

import fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from app.converters.positioned import (
    _add_behind_picture,
    _add_textbox,
    _escape_xml,
    _iter_text_lines,
    _rgb_from_span,
    _sample_fill_color,
    _zero_paragraph_spacing,
    pdf_to_docx_positioned,
)
from tests.support import add_page, write_pdf


class _Page:
    def __init__(self, blocks):
        self._blocks = blocks

    def get_text(self, kind="dict", flags=0):
        return {"blocks": self._blocks}


def test_color_escape_and_fill_sampling():
    assert _rgb_from_span({}) == (0, 0, 0)
    assert _rgb_from_span({"color": None}) == (0, 0, 0)
    assert _rgb_from_span({"color": 0xFF8001}) == (255, 128, 1)
    assert _rgb_from_span({"color": "red"}) == (0, 0, 0)
    assert _escape_xml('A & B <tag> "q"') == "A &amp; B &lt;tag&gt; &quot;q&quot;"

    image = Image.new("RGB", (30, 20), (10, 20, 30))
    assert _sample_fill_color(image, (4, 4, 18, 16)) == (10, 20, 30)
    assert _sample_fill_color(image, (5, 5, 5, 12)) == (255, 255, 255)
    assert _sample_fill_color(image, (100, 100, 120, 140)) == (255, 255, 255)
    edge = _sample_fill_color(image, (-4, -4, 8, 6))
    assert edge == (10, 20, 30)


def test_iter_text_lines_skips_incomplete_blocks_and_styles_spans():
    page = _Page(
        [
            {"type": 1, "lines": []},
            {"type": 0, "lines": []},
            {"type": 0, "lines": [{"spans": []}]},
            {"type": 0, "lines": [{"spans": [{"text": "   ", "size": 10, "bbox": (1, 1, 2, 2)}]}]},
            {"type": 0, "lines": [{"spans": [{"text": "nobox", "size": 10}]}]},
            {
                "type": 0,
                "lines": [
                    {
                        "bbox": (10, 10, 40, 20),
                        "spans": [
                            {"text": "Hi", "size": 12, "flags": 16, "font": "Helvetica", "color": 255},
                            {"text": "!", "size": 8, "flags": 0, "font": "Helvetica", "color": 0},
                        ],
                    }
                ],
            },
            {
                "type": 0,
                "lines": [
                    {
                        "spans": [
                            {
                                "text": "BoldName",
                                "size": 14,
                                "flags": 0,
                                "font": "Arial-Bold",
                                "color": None,
                                "bbox": (0, 30, 20, 44),
                            }
                        ]
                    }
                ],
            },
        ]
    )
    lines = _iter_text_lines(page)
    assert [line["text"] for line in lines] == ["Hi!", "BoldName"]
    assert lines[0]["bold"] is True
    assert lines[0]["rgb"] == (0, 0, 255)
    assert lines[0]["bbox"] == (10.0, 10.0, 40.0, 20.0)
    assert lines[1]["bold"] is True
    assert lines[1]["rgb"] == (0, 0, 0)
    assert lines[1]["bbox"] == (0.0, 30.0, 20.0, 44.0)


def test_textbox_and_spacing_markup():
    document = Document()
    paragraph = document.add_paragraph("seed")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "200")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "100")
    align = OxmlElement("w:jc")
    align.set(qn("w:val"), "center")
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(spacing)
    p_pr.append(indent)
    p_pr.append(align)

    _zero_paragraph_spacing(paragraph)
    tags = [child.tag for child in paragraph._p.get_or_add_pPr()]
    assert qn("w:ind") not in tags
    assert qn("w:jc") in tags
    kept = paragraph._p.find(qn("w:pPr")).find(qn("w:spacing"))
    assert kept.get(qn("w:before")) == "0"
    assert kept.get(qn("w:after")) == "0"

    _add_textbox(
        paragraph,
        text='A & B <tag> "q"',
        left=-20,
        top=-10,
        width=1,
        height=1,
        font_size_pt=1,
        rgb=(1, 2, 3),
        bold=True,
    )
    _add_textbox(
        paragraph,
        text="plain",
        left=1000,
        top=1000,
        width=20000,
        height=20000,
        font_size_pt=12,
        rgb=(0, 0, 0),
        bold=False,
    )
    xml = paragraph._p.xml
    assert 'A &amp; B &lt;tag&gt; "q"' in xml
    assert "<w:b/>" in xml or "<w:b />" in xml
    assert 'w:val="10"' in xml
    assert 'w:val="24"' in xml
    assert ">plain<" in xml


def test_behind_picture_returns_when_drawing_is_missing():
    document = Document()
    paragraph = document.paragraphs[0] if document.paragraphs else document.add_paragraph()

    class _DocPr:
        def __init__(self, values):
            self.values = values

        def get(self, key, default=None):
            return self.values.get(key, default)

    class _Run:
        def __init__(self, doc_pr):
            self._inline = SimpleNamespace(
                docPr=doc_pr,
                graphic=SimpleNamespace(
                    graphicData=SimpleNamespace(
                        pic=SimpleNamespace(
                            blipFill=SimpleNamespace(blip=SimpleNamespace(get=lambda key: "rId1"))
                        )
                    )
                ),
            )

        def add_picture(self, *_args, **_kwargs):
            return SimpleNamespace(_inline=self._inline)

        class _r:
            @staticmethod
            def find(_tag):
                return None

    paragraph.add_run = lambda: _Run(_DocPr({"id": "", "name": ""}))
    buffer = BytesIO()
    Image.new("RGB", (8, 8), "white").save(buffer, format="PNG")
    buffer.seek(0)
    assert _add_behind_picture(paragraph, buffer, 100, 80) is None

    paragraph.add_run = lambda: _Run(_DocPr({"id": "7", "name": "Backdrop"}))
    buffer.seek(0)
    assert _add_behind_picture(paragraph, buffer, 100, 80) is None


def test_positioned_conversion_preserves_text_and_extra_pages(tmp_path, monkeypatch):
    def populate(doc):
        page = add_page(doc, "Hello & friends")
        page.insert_text((72, 120), "Bold", fontname="hebo", fontsize=16, color=(1, 0, 0))
        page.insert_text((72, 160), "   ")
        add_page(doc, "Second page")

    src = write_pdf(tmp_path / "in.pdf", populate)
    dst = tmp_path / "out.docx"
    assert pdf_to_docx_positioned(src, dst) == dst
    with zipfile.ZipFile(dst) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert "Hello &amp; friends" in xml
    assert "Second page" in xml
    assert media
    document = Document(str(dst))
    assert len(document.sections) == 2

    original = pdf_to_docx_positioned.__globals__["Document"]

    def seeded():
        document = original()
        document.add_paragraph("seed")
        return document

    monkeypatch.setattr("app.converters.positioned.Document", seeded)
    seeded = tmp_path / "seed.docx"
    single = write_pdf(tmp_path / "one.pdf", lambda doc: add_page(doc, "Only"))
    pdf_to_docx_positioned(single, seeded)
    with zipfile.ZipFile(seeded) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        assert any(name.startswith("word/media/") for name in archive.namelist())
    assert "Only" in xml


def test_real_page_lines_skip_whitespace(tmp_path):
    src = write_pdf(
        tmp_path / "lines.pdf",
        lambda doc: add_page(doc, "Hello"),
    )
    with fitz.open(src) as pdf:
        pdf[0].insert_text((72, 140), "   ")
        lines = _iter_text_lines(pdf[0])
    assert [line["text"].strip() for line in lines] == ["Hello"]
    assert lines[0]["bold"] is False
