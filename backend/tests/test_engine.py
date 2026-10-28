"""Format routing and the concrete converters."""

from __future__ import annotations

import zipfile
from pathlib import Path
from types import SimpleNamespace

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu
from PIL import Image

from app.converters import engine
from app.converters.engine import (
    SUPPORTED_TARGETS,
    _find_soffice,
    _zero_paragraph_spacing,
    convert_file,
    convert_image,
    convert_with_libreoffice,
    docx_to_html,
    docx_to_txt,
    images_to_pdf,
    libreoffice_available,
    pdf_to_docx,
    pdf_to_docx_editable,
    pdf_to_docx_flow,
    pdf_to_docx_visual,
    pdf_to_html,
    pdf_to_images,
    pdf_to_text,
    text_to_docx,
    text_to_pdf,
)
from tests.support import add_page, solid_image, write_pdf

_PT_TO_EMU = 12700


def _run_result(code=0, stdout="", stderr=""):
    return SimpleNamespace(returncode=code, stdout=stdout, stderr=stderr)


def _install_soffice(monkeypatch, runner):
    monkeypatch.setattr(engine, "_find_soffice", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(engine.subprocess, "run", runner)


def _no_soffice(monkeypatch):
    monkeypatch.setattr(engine, "_find_soffice", lambda: None)


def test_find_soffice_and_availability(monkeypatch):
    monkeypatch.setattr(
        engine.shutil, "which", lambda name: "/bin/soffice" if name == "soffice" else None
    )
    assert _find_soffice() == "/bin/soffice"

    monkeypatch.setattr(
        engine.shutil,
        "which",
        lambda name: "/bin/libreoffice" if name == "libreoffice" else None,
    )
    assert _find_soffice() == "/bin/libreoffice"
    assert libreoffice_available() is True

    monkeypatch.setattr(engine.shutil, "which", lambda name: None)
    real_exists = Path.exists
    hit = {"path": "/usr/lib/libreoffice/program/soffice"}

    def exists(self):
        if str(self).endswith("soffice"):
            return str(self) == hit["path"]
        return real_exists(self)

    monkeypatch.setattr(Path, "exists", exists)
    assert _find_soffice() == "/usr/lib/libreoffice/program/soffice"
    hit["path"] = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    assert _find_soffice() == hit["path"]
    hit["path"] = "/usr/bin/soffice"
    assert _find_soffice() == hit["path"]
    hit["path"] = ""
    assert _find_soffice() is None
    assert libreoffice_available() is False
    assert "docx" in SUPPORTED_TARGETS


def test_convert_with_libreoffice(monkeypatch, tmp_path):
    src = tmp_path / "note.docx"
    src.write_text("doc")
    out = tmp_path / "out"
    out.mkdir()
    _no_soffice(monkeypatch)
    with pytest.raises(RuntimeError, match="LibreOffice 未安裝"):
        convert_with_libreoffice(src, "pdf", out)

    def succeed(cmd, capture_output, text, timeout):
        assert cmd[0] == "/usr/bin/soffice"
        assert "--headless" in cmd
        assert capture_output is True and text is True and timeout == 180
        produced = out / "note.pdf"
        produced.write_bytes(b"%PDF")
        return _run_result()

    _install_soffice(monkeypatch, succeed)
    assert convert_with_libreoffice(src, "pdf", out).read_bytes() == b"%PDF"
    with pytest.raises(ValueError, match="webp"):
        convert_with_libreoffice(src, "webp", out)

    def missing_but_sibling(cmd, **kwargs):
        (out / "note.odt").write_text("alt")
        return _run_result()

    monkeypatch.setattr(engine.subprocess, "run", missing_but_sibling)
    (out / "note.pdf").unlink()
    assert convert_with_libreoffice(src, "pdf", out).name == "note.odt"

    def nothing(cmd, **kwargs):
        return _run_result()

    monkeypatch.setattr(engine.subprocess, "run", nothing)
    (out / "note.odt").unlink()
    with pytest.raises(RuntimeError, match="未產生輸出檔案"):
        convert_with_libreoffice(src, "pdf", out)

    monkeypatch.setattr(
        engine.subprocess, "run", lambda *a, **k: _run_result(1, stdout="", stderr="bad-err")
    )
    with pytest.raises(RuntimeError, match="bad-err"):
        convert_with_libreoffice(src, "docx", out)
    monkeypatch.setattr(
        engine.subprocess, "run", lambda *a, **k: _run_result(1, stdout="bad-out", stderr="")
    )
    with pytest.raises(RuntimeError, match="bad-out"):
        convert_with_libreoffice(src, "docx", out)
    monkeypatch.setattr(
        engine.subprocess, "run", lambda *a, **k: _run_result(1, stdout="", stderr="")
    )
    with pytest.raises(RuntimeError, match="LibreOffice 轉換失敗:"):
        convert_with_libreoffice(src, "docx", out)


def test_zero_paragraph_spacing_removes_only_spacing_and_indent():
    document = Document()
    paragraph = document.add_paragraph("hello")
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "120")
    indent = OxmlElement("w:ind")
    align = OxmlElement("w:jc")
    align.set(qn("w:val"), "right")
    p_pr.extend([spacing, indent, align])
    _zero_paragraph_spacing(paragraph)
    tags = [child.tag for child in paragraph._p.find(qn("w:pPr"))]
    assert tags.count(qn("w:spacing")) == 1
    assert qn("w:ind") not in tags
    assert qn("w:jc") in tags
    assert paragraph._p.find(qn("w:pPr")).find(qn("w:spacing")).get(qn("w:before")) == "0"


def test_visual_docx_page_geometry(tmp_path, monkeypatch):
    src = write_pdf(
        tmp_path / "two.pdf",
        lambda doc: (add_page(doc, "one"), add_page(doc, "two")),
    )
    dst = tmp_path / "two.docx"
    assert pdf_to_docx_visual(src, dst) == dst
    with fitz.open(src) as pdf:
        width = float(pdf[0].rect.width)
    document = Document(str(dst))
    assert len(document.sections) == 2
    assert document.sections[0].page_width == Emu(int(width * _PT_TO_EMU))
    with zipfile.ZipFile(dst) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert len(media) == 2

    original_document = engine.Document

    def seeded():
        document = original_document()
        document.add_paragraph("seed")
        return document

    monkeypatch.setattr(engine, "Document", seeded)
    single = write_pdf(tmp_path / "one.pdf", lambda doc: add_page(doc, "only"))
    seeded = tmp_path / "seed.docx"
    pdf_to_docx_visual(single, seeded)
    with zipfile.ZipFile(seeded) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())


def test_flow_and_editable_wrappers(monkeypatch, tmp_path):
    src = tmp_path / "in.pdf"
    src.write_bytes(b"%PDF")
    created = []

    class FakeConverter:
        def __init__(self, path):
            self.path = path
            self.closed = False
            self.kwargs = None
            created.append(self)

        def convert(self, dst, **kwargs):
            self.kwargs = kwargs
            Path(dst).write_text("flow")

        def close(self):
            self.closed = True

    monkeypatch.setattr(engine, "Converter", FakeConverter)
    dst = tmp_path / "flow.docx"
    assert pdf_to_docx_flow(src, dst) == dst
    assert dst.read_text() == "flow"
    assert created[0].closed is True
    assert created[0].kwargs["parse_stream_table"] is False
    assert created[0].kwargs["delete_end_line_hyphen"] is True

    class Boom(FakeConverter):
        def convert(self, dst, **kwargs):
            raise RuntimeError("flow failed")

    monkeypatch.setattr(engine, "Converter", Boom)
    with pytest.raises(RuntimeError, match="flow failed"):
        pdf_to_docx_flow(src, tmp_path / "bad.docx")
    assert created[-1].closed is True

    def positioned(src_path, dst_path, **kwargs):
        dst_path.write_text("positioned")
        return dst_path

    monkeypatch.setattr("app.converters.positioned.pdf_to_docx_positioned", positioned)
    editable = tmp_path / "edit.docx"
    assert pdf_to_docx_editable(src, editable).read_text() == "positioned"

    def explode(*args, **kwargs):
        raise OSError("nope")

    monkeypatch.setattr("app.converters.positioned.pdf_to_docx_positioned", explode)
    monkeypatch.setattr(engine, "Converter", FakeConverter)
    fallback = tmp_path / "fallback.docx"
    assert pdf_to_docx_editable(src, fallback).read_text() == "flow"


def test_pdf_to_docx_delegates_to_pipeline(monkeypatch, tmp_path):
    src = tmp_path / "in.pdf"
    src.write_bytes(b"%PDF")
    seen = {}

    def fake_pipeline(src_path, dst, base, mode="auto"):
        seen["base"] = base
        seen["mode"] = mode
        dst.write_text("piped")
        return SimpleNamespace(path=dst)

    monkeypatch.setattr(
        "app.converters.pipeline.convert_pdf_to_docx_pipeline", fake_pipeline
    )
    dst = tmp_path / "nested" / "out.docx"
    dst.parent.mkdir()
    assert pdf_to_docx(src, dst, mode=None).read_text() == "piped"
    assert seen["mode"] == "auto"
    assert seen["base"] == dst.parent

    work = tmp_path / "work"
    work.mkdir()
    pdf_to_docx(src, dst, mode=" Editable ", work_dir=work)
    assert seen["mode"] == "editable"
    assert seen["base"] == work
    with pytest.raises(ValueError, match="layout_mode"):
        pdf_to_docx(src, dst, mode="sideways")


def test_pdf_text_html_and_images(tmp_path):
    src = write_pdf(
        tmp_path / "pages.pdf",
        lambda doc: (add_page(doc, "Alpha"), add_page(doc, "Beta")),
    )
    text = pdf_to_text(src, tmp_path / "pages.txt")
    assert "Alpha" in text.read_text(encoding="utf-8")
    assert "\n\n" in text.read_text(encoding="utf-8")
    html = pdf_to_html(src, tmp_path / "pages.html").read_text(encoding="utf-8")
    assert "Page 1" in html and "Page 2" in html and "Beta" in html

    images = pdf_to_images(src, tmp_path, fmt="png")
    assert [path.name for path in images] == ["pages_p1.png", "pages_p2.png"]
    assert Image.open(images[0]).format == "PNG"
    jpgs = pdf_to_images(write_pdf(tmp_path / "one.pdf", lambda doc: add_page(doc, "Solo")), tmp_path, fmt="jpg")
    assert jpgs[0].suffix == ".jpg"


def test_docx_and_text_and_image_helpers(tmp_path):
    document = Document()
    document.add_paragraph("Para")
    document.add_paragraph("")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Left"
    table.rows[0].cells[1].text = "Right"
    docx = tmp_path / "File.DOCX"
    document.save(docx)
    assert "Para" in docx_to_txt(docx, tmp_path / "out.txt").read_text(encoding="utf-8")
    assert "Left\tRight" in (tmp_path / "out.txt").read_text(encoding="utf-8")
    html = docx_to_html(docx, tmp_path / "out.html").read_text(encoding="utf-8")
    assert "Para" in html and "<html" in html

    note = tmp_path / "note.txt"
    note.write_bytes(b"line1\nline2\n\xff\n")
    rebuilt = text_to_docx(note, tmp_path / "note.docx")
    paragraphs = [p.text for p in Document(str(rebuilt)).paragraphs]
    assert "line1" in paragraphs and "line2" in paragraphs
    pdf = text_to_pdf(note, tmp_path / "note.pdf")
    assert "line1" in pdf_to_text(pdf, tmp_path / "back.txt").read_text(encoding="utf-8")

    rgba = solid_image(tmp_path / "rgba.png", color=(255, 0, 0, 128), mode="RGBA")
    palette = solid_image(tmp_path / "pal.png", mode="P")
    rgb = solid_image(tmp_path / "rgb.png", color=(0, 128, 0))
    pdf_path = images_to_pdf([rgba, palette, rgb], tmp_path / "images.pdf")
    with fitz.open(pdf_path) as rendered:
        assert rendered.page_count == 3
    jpeg = convert_image(rgba, tmp_path / "out.jpeg")
    assert jpeg.read_bytes()[:3] == b"\xff\xd8\xff"
    assert convert_image(rgb, tmp_path / "plain.jpg").read_bytes()[:3] == b"\xff\xd8\xff"
    assert Image.open(convert_image(palette, tmp_path / "pal.jpg")).mode == "RGB"
    assert Image.open(convert_image(rgb, tmp_path / "out.png")).format == "PNG"
    assert convert_image(rgb, tmp_path / "out.webp").stat().st_size > 0
    assert convert_image(rgb, tmp_path / "out.gif").read_bytes()[:3] == b"GIF"


def test_convert_file_routes_local_formats(monkeypatch, tmp_path):
    _no_soffice(monkeypatch)
    work = tmp_path / "work"
    work.mkdir()
    pdf = write_pdf(
        tmp_path / "Doc.PDF",
        lambda doc: (add_page(doc, "Alpha"), add_page(doc, "Beta")),
    )
    single = write_pdf(tmp_path / "single.pdf", lambda doc: add_page(doc, "Solo"))

    copied = convert_file(_write(tmp_path / "same.txt", "same"), "txt", work)
    assert copied.read_text(encoding="utf-8") == "same"
    assert copied.parent == work

    text = convert_file(single, "txt", work)
    assert "Solo" in text.read_text(encoding="utf-8")
    html = convert_file(pdf, "html", work)
    assert "Alpha" in html.read_text(encoding="utf-8")
    png = convert_file(single, "png", work)
    assert png.suffix == ".png" and png.name == "single.png"
    jpg = convert_file(single, ".JPG", work)
    assert jpg.suffix == ".jpg"
    zipped = convert_file(pdf, "png", work)
    assert zipped.suffix == ".zip"
    with zipfile.ZipFile(zipped) as archive:
        assert len(archive.namelist()) == 2

    seen = {}

    def fake_pipeline(src, dst, base, mode="auto"):
        seen["mode"] = mode
        seen["base"] = base
        dst.write_bytes(b"docx")
        return SimpleNamespace(path=dst)

    monkeypatch.setattr("app.converters.pipeline.convert_pdf_to_docx_pipeline", fake_pipeline)
    produced = convert_file(single, "DOCX", work, layout_mode="hybrid")
    assert produced.read_bytes() == b"docx"
    assert seen["mode"] == "hybrid"

    docx = tmp_path / "letter.DOCX"
    document = Document()
    document.add_paragraph("Body")
    document.save(docx)
    assert "Body" in convert_file(docx, "html", work).read_text(encoding="utf-8")
    assert "Body" in convert_file(docx, "txt", work).read_text(encoding="utf-8")

    image = solid_image(tmp_path / "shot.png", mode="RGBA", color=(1, 2, 3, 4))
    assert convert_file(image, "JPEG", work).suffix == ".jpg"
    assert convert_file(image, "webp", work).suffix == ".webp"
    assert convert_file(image, "pdf", work).read_bytes().startswith(b"%PDF")
    pictured = convert_file(image, "docx", work)
    assert len(Document(str(pictured)).inline_shapes) == 1
    gif = solid_image(tmp_path / "anim.gif", mode="RGB")
    assert convert_file(gif, "png", work).suffix == ".png"

    markdown = _write(tmp_path / "note.md", "hello md")
    assert "hello md" in convert_file(markdown, "html", work).read_text(encoding="utf-8")
    assert convert_file(markdown, "txt", work).read_text(encoding="utf-8") == "hello md"
    assert "hello md" in pdf_to_text(convert_file(markdown, "pdf", work), work / "md-back.txt").read_text()
    assert any(p.text == "hello md" for p in Document(str(convert_file(markdown, "docx", work))).paragraphs)
    log = _write(tmp_path / "app.log", "log line")
    assert convert_file(log, "html", work).read_text(encoding="utf-8").count("log line") == 1
    csv = _write(tmp_path / "rows.csv", "a,b")
    assert convert_file(csv, "txt", work).read_text(encoding="utf-8") == "a,b"
    broken = tmp_path / "bad.txt"
    broken.write_bytes(b"ok\xff")
    repaired = convert_file(broken, "docx", work)
    assert any("ok" in paragraph.text for paragraph in Document(str(repaired)).paragraphs)
    assert "ok" in convert_file(broken, "html", work).read_text(encoding="utf-8")

    page = _write(tmp_path / "page.html", "<p>Hello <b>there</b></p>")
    assert convert_file(page, "txt", work).read_text(encoding="utf-8") == "Hello there"
    htm = _write(tmp_path / "page.htm", "<div>Htm</div>")
    assert "Htm" in convert_file(htm, "txt", work).read_text(encoding="utf-8")


def _write(path: Path, text: str) -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def test_convert_file_libreoffice_and_errors(monkeypatch, tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    _no_soffice(monkeypatch)

    def fake_pdf_to_docx(src, dst, mode="auto", work_dir=None):
        dst.write_bytes(b"interim")
        return dst

    monkeypatch.setattr(engine, "pdf_to_docx", fake_pdf_to_docx)
    pdf = write_pdf(tmp_path / "in.pdf", lambda doc: add_page(doc, "Text"))
    docx = tmp_path / "in.docx"
    Document().save(docx)
    binary = tmp_path / "old.doc"
    binary.write_bytes(b"doc")
    html = _write(tmp_path / "in.html", "<p>Hi</p>")
    pptx = tmp_path / "deck.pptx"
    pptx.write_bytes(b"pptx")
    unknown = tmp_path / "data.xyz"
    unknown.write_bytes(b"xyz")

    with pytest.raises(ValueError, match="不支援的目標格式"):
        convert_file(pdf, "exe", work)
    with pytest.raises(RuntimeError, match="PDF→DOC"):
        convert_file(pdf, "doc", work)
    with pytest.raises(RuntimeError, match="LibreOffice"):
        convert_file(docx, "pdf", work)
    with pytest.raises(RuntimeError, match="LibreOffice"):
        convert_file(binary, "html", work)
    with pytest.raises(ValueError, match="安裝 LibreOffice"):
        convert_file(unknown, "txt", work)
    with pytest.raises(ValueError, match="目前不支援"):
        convert_file(html, "docx", work)
    with pytest.raises(ValueError, match="目前不支援"):
        convert_file(pdf, "webp", work)
    with pytest.raises(ValueError, match="目前不支援"):
        convert_file(_write(tmp_path / "plain.txt", "x"), "png", work)
    with pytest.raises(ValueError, match="目前不支援"):
        convert_file(solid_image(tmp_path / "lone.png"), "txt", work)

    calls = []

    def runner(cmd, capture_output, text, timeout):
        target = cmd[cmd.index("--convert-to") + 1]
        src = Path(cmd[-1])
        produced = Path(cmd[cmd.index("--outdir") + 1]) / f"{src.stem}.{target}"
        produced.write_bytes(b"lo-" + target.encode())
        calls.append((src.name, target))
        return _run_result()

    _install_soffice(monkeypatch, runner)
    monkeypatch.setattr(
        "app.converters.pipeline.convert_pdf_to_docx_pipeline",
        lambda src, dst, base, mode="auto": dst.write_bytes(b"interim") or SimpleNamespace(path=dst),
    )
    assert convert_file(pdf, "doc", work, layout_mode="visual").read_bytes() == b"lo-doc"
    assert ("in.docx", "doc") in calls or calls[-1][1] == "doc"

    assert convert_file(docx, "pdf", work).read_bytes() == b"lo-pdf"
    assert convert_file(docx, "odt", work).read_bytes() == b"lo-odt"
    assert convert_file(docx, "rtf", work).read_bytes() == b"lo-rtf"
    assert convert_file(docx, ".DOC", work).read_bytes() == b"lo-doc"
    assert convert_file(binary, "docx", work).read_bytes() == b"lo-docx"
    assert convert_file(html, "pdf", work).read_bytes() == b"lo-pdf"
    assert convert_file(html, "docx", work).read_bytes() == b"lo-docx"
    assert convert_file(pptx, "pdf", work).read_bytes() == b"lo-pdf"
    assert convert_file(pdf, "odt", work).read_bytes() == b"lo-odt"
    assert convert_file(unknown, "pdf", work).read_bytes() == b"lo-pdf"
    with pytest.raises(ValueError, match="目前不支援") as unsupported:
        convert_file(html, "png", work)
    assert "安裝 LibreOffice" not in str(unsupported.value)

    with pytest.raises(ValueError, match="LibreOffice 不支援輸出格式"):
        convert_file(pptx, "webp", work)
    caught = None
    try:
        convert_file(unknown, "txt", work)
    except ValueError as exc:
        caught = str(exc)
    assert caught is not None
    assert "目前不支援" in caught
    assert "安裝 LibreOffice" not in caught

    _no_soffice(monkeypatch)
    # html → txt still strips tags when LibreOffice is absent; html → pdf does not.
    assert "Hi" in convert_file(html, "txt", work).read_text(encoding="utf-8")
    with pytest.raises(ValueError, match="安裝 LibreOffice"):
        convert_file(html, "pdf", work)
