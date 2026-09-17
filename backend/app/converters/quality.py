"""Conversion Quality Score — explainable metrics for PDF→DOCX.

Metrics (0–100%):
  - layout_similarity
  - text_preservation
  - image_preservation
  - element_alignment
  - page_structure

Overall is a layout-aware weighted average of the five.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from typing import Any

import fitz
import numpy as np
from docx import Document
from PIL import Image


_SPACE_LETTER = re.compile(
    r"(?:\b(?:[A-Za-z0-9]\s+){3,}[A-Za-z0-9]\b)|(?:(?:[\u4e00-\u9fff]\s+){3,}[\u4e00-\u9fff])"
)

# If layout similarity falls below this after an editable attempt, prefer visual.
LAYOUT_RETRY_THRESHOLD = 70.0
OVERALL_RETRY_THRESHOLD = 65.0


def _clamp(value: float) -> float:
    return max(0.0, min(100.0, value))


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text.lower()).strip()


def _pdf_text(path: Path) -> str:
    doc = fitz.open(path)
    try:
        return "\n".join(page.get_text("text") or "" for page in doc)
    finally:
        doc.close()


def _docx_text(path: Path) -> str:
    document = Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _docx_stats(path: Path) -> dict[str, int]:
    document = Document(str(path))
    return {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "sections": len(document.sections),
        "inline_shapes": len(document.inline_shapes),
    }


def _spaced_out_ratio(text: str) -> float:
    if not text.strip():
        return 0.0
    hits = _SPACE_LETTER.findall(text)
    hit_len = sum(len(h) for h in hits)
    return min(1.0, hit_len / max(1, len(text)))


def _text_overlap(a: str, b: str) -> float:
    ta = set(_normalize(a).split())
    tb = set(_normalize(b).split())
    if not ta and not tb:
        return 1.0
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)


def _pdf_image_count(path: Path) -> int:
    doc = fitz.open(path)
    try:
        total = 0
        for page in doc:
            total += len(page.get_images(full=True))
        return total
    finally:
        doc.close()


def _docx_media_images(path: Path) -> list[Image.Image]:
    """Extract embedded images from a DOCX package."""
    images: list[Image.Image] = []
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if not name.startswith("word/media/"):
                continue
            lower = name.lower()
            if not lower.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp")):
                continue
            try:
                raw = zf.read(name)
                img = Image.open(io.BytesIO(raw)).convert("RGB")
                images.append(img)
            except Exception:  # noqa: BLE001
                continue
    return images


def _render_pdf_pages(path: Path, *, max_pages: int = 8, scale: float = 0.35) -> list[Image.Image]:
    doc = fitz.open(path)
    pages: list[Image.Image] = []
    try:
        matrix = fitz.Matrix(scale, scale)
        for i, page in enumerate(doc):
            if i >= max_pages:
                break
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            pages.append(img)
    finally:
        doc.close()
    return pages


def _image_similarity(a: Image.Image, b: Image.Image) -> float:
    """Normalized similarity in [0, 1] via mean absolute difference."""
    w = min(a.width, b.width, 480)
    h = min(a.height, b.height, 640)
    if w < 8 or h < 8:
        return 0.0
    aa = np.asarray(a.resize((w, h), Image.Resampling.BILINEAR), dtype=np.float32)
    bb = np.asarray(b.resize((w, h), Image.Resampling.BILINEAR), dtype=np.float32)
    mae = float(np.mean(np.abs(aa - bb)))
    return max(0.0, 1.0 - mae / 255.0)


def _aspect_ratio_score(w: float, h: float, tw: float, th: float) -> float:
    if min(w, h, tw, th) <= 0:
        return 0.0
    r1 = w / h
    r2 = tw / th
    diff = abs(r1 - r2) / max(r1, r2)
    return max(0.0, 1.0 - diff * 2.0)


def _layout_similarity(
    pdf_path: Path,
    docx_path: Path,
    *,
    mode_used: str,
    layout: dict[str, Any],
    spaced: float,
    overlap: float,
    stats: dict[str, int],
) -> tuple[float, dict[str, Any]]:
    """Explainable layout similarity (%)."""
    detail: dict[str, Any] = {"method": mode_used}

    if mode_used == "visual":
        pdf_pages = _render_pdf_pages(pdf_path)
        media = _docx_media_images(docx_path)
        detail["pdf_pages_compared"] = len(pdf_pages)
        detail["docx_media_images"] = len(media)
        if not pdf_pages:
            return 90.0, {**detail, "note": "empty_pdf"}
        sims: list[float] = []
        aspects: list[float] = []
        for i, page_img in enumerate(pdf_pages):
            if i < len(media):
                sims.append(_image_similarity(page_img, media[i]))
                aspects.append(
                    _aspect_ratio_score(
                        page_img.width, page_img.height, media[i].width, media[i].height
                    )
                )
        page_sim = float(np.mean(sims)) if sims else 0.85
        aspect = float(np.mean(aspects)) if aspects else 0.9
        # Visual path preserves pixels by construction; image compare validates packaging
        score = 100.0 * (0.75 * page_sim + 0.25 * aspect)
        # Floor: visual should remain strong even if resize noise
        score = max(score, 88.0 if sims else 92.0)
        detail.update(
            {
                "page_image_similarity": round(page_sim, 3),
                "aspect_ratio_score": round(aspect, 3),
            }
        )
        return _clamp(score), detail

    # Editable / heuristic path — no DOCX rasterizer without LibreOffice
    detail["method"] = "heuristic_editable"
    base = 42.0 + 48.0 * overlap - 45.0 * spaced
    if layout.get("layout_type") in {"designed", "complex", "scanned"}:
        base *= 0.5
        detail["complex_layout_penalty"] = True
    pages = max(1, int(layout.get("pages") or 1))
    if stats["sections"] > max(3, pages * 2):
        base -= 15
        detail["section_explosion"] = stats["sections"]
    detail["text_overlap"] = round(overlap, 3)
    detail["spaced_out_ratio"] = round(spaced, 3)
    return _clamp(base), detail


def _text_preservation(
    *,
    mode_used: str,
    pdf_chars: int,
    docx_chars: int,
    overlap: float,
    spaced: float,
    is_scanned: bool,
) -> tuple[float, dict[str, Any]]:
    if mode_used == "visual":
        # Glyphs preserved as page imagery (not extractable text)
        score = 96.0 if pdf_chars > 0 or is_scanned else 80.0
        return score, {
            "method": "visual_raster_preservation",
            "note": "Text is preserved as page pixels, not as editable runs.",
            "pdf_chars": pdf_chars,
        }

    coverage = min(1.0, docx_chars / max(1, pdf_chars)) if pdf_chars else 1.0
    score = 100.0 * (0.7 * overlap + 0.3 * coverage) - 40.0 * spaced
    return _clamp(score), {
        "method": "token_overlap_and_coverage",
        "text_overlap": round(overlap, 3),
        "char_coverage": round(coverage, 3),
        "spaced_out_ratio": round(spaced, 3),
        "pdf_chars": pdf_chars,
        "docx_chars": docx_chars,
    }


def _image_preservation(pdf_path: Path, docx_path: Path, *, mode_used: str) -> tuple[float, dict[str, Any]]:
    pdf_images = _pdf_image_count(pdf_path)
    media = _docx_media_images(docx_path)
    docx_images = len(media)
    detail = {
        "pdf_images": pdf_images,
        "docx_images": docx_images,
        "mode_used": mode_used,
    }
    if mode_used == "visual":
        # One rendered page image per page counts as preservation of visual content
        doc = fitz.open(pdf_path)
        try:
            pdf_pages = doc.page_count
        finally:
            doc.close()
        detail["pdf_pages"] = pdf_pages
        if docx_images >= max(1, pdf_pages):
            return 100.0, detail
        if docx_images > 0:
            return _clamp(100.0 * docx_images / max(1, pdf_pages)), detail
        return 40.0, detail

    if pdf_images == 0:
        return 100.0, {**detail, "note": "source_has_no_images"}
    ratio = min(1.0, docx_images / pdf_images)
    return _clamp(100.0 * ratio), detail


def _element_alignment(
    *,
    mode_used: str,
    spaced: float,
    stats: dict[str, int],
    layout: dict[str, Any],
    layout_sim_detail: dict[str, Any],
) -> tuple[float, dict[str, Any]]:
    if mode_used == "visual":
        aspect = float(layout_sim_detail.get("aspect_ratio_score") or 0.9)
        score = 100.0 * (0.55 + 0.45 * aspect)
        return _clamp(score), {
            "method": "page_image_aspect_alignment",
            "aspect_ratio_score": round(aspect, 3),
        }

    score = 90.0 - 55.0 * spaced
    pages = max(1, int(layout.get("pages") or 1))
    if stats["sections"] > max(3, pages * 2):
        score -= 18
    if layout.get("columns_estimate", 1) >= 2:
        score -= 12
    return _clamp(score), {
        "method": "editable_flow_alignment",
        "spaced_out_ratio": round(spaced, 3),
        "sections": stats["sections"],
        "columns_estimate": layout.get("columns_estimate"),
    }


def _page_structure(
    pdf_path: Path,
    *,
    mode_used: str,
    stats: dict[str, int],
) -> tuple[float, dict[str, Any]]:
    doc = fitz.open(pdf_path)
    try:
        pdf_pages = doc.page_count
    finally:
        doc.close()

    sections = stats["sections"]
    detail = {
        "pdf_pages": pdf_pages,
        "docx_sections": sections,
        "docx_paragraphs": stats["paragraphs"],
        "docx_tables": stats["tables"],
    }

    if mode_used == "visual":
        # Expect ~1 section (or page) image per PDF page
        if sections == pdf_pages or stats["inline_shapes"] >= pdf_pages:
            score = 98.0 if sections == pdf_pages else 94.0
        else:
            score = 100.0 * min(sections, pdf_pages) / max(1, pdf_pages)
        return _clamp(score), detail

    # Editable: pages may map to flowing sections; penalize explosion or emptiness
    if stats["paragraphs"] + stats["tables"] == 0:
        return 20.0, detail
    if sections > max(4, pdf_pages * 3):
        return _clamp(55.0), {**detail, "note": "section_explosion"}
    if sections >= 1:
        return 92.0, detail
    return 70.0, detail


def score_conversion(
    pdf_path: Path,
    docx_path: Path,
    *,
    mode_used: str,
    layout: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Build explainable Conversion Quality Score."""
    layout = layout or {}
    pdf_txt = _pdf_text(pdf_path)
    docx_txt = _docx_text(docx_path)
    stats = _docx_stats(docx_path)
    spaced = _spaced_out_ratio(docx_txt)
    overlap = _text_overlap(pdf_txt, docx_txt)
    pdf_chars = len(pdf_txt.strip())
    docx_chars = len(docx_txt.strip())

    layout_similarity, layout_detail = _layout_similarity(
        pdf_path,
        docx_path,
        mode_used=mode_used,
        layout=layout,
        spaced=spaced,
        overlap=overlap,
        stats=stats,
    )
    text_preservation, text_detail = _text_preservation(
        mode_used=mode_used,
        pdf_chars=pdf_chars,
        docx_chars=docx_chars,
        overlap=overlap,
        spaced=spaced,
        is_scanned=bool(layout.get("is_scanned")),
    )
    image_preservation, image_detail = _image_preservation(
        pdf_path, docx_path, mode_used=mode_used
    )
    element_alignment, align_detail = _element_alignment(
        mode_used=mode_used,
        spaced=spaced,
        stats=stats,
        layout=layout,
        layout_sim_detail=layout_detail,
    )
    page_structure, structure_detail = _page_structure(
        pdf_path, mode_used=mode_used, stats=stats
    )

    # Weights: designed/complex emphasize layout; simple emphasizes text
    if layout.get("layout_type") in {"designed", "complex", "scanned"}:
        weights = {
            "layout_similarity": 0.35,
            "text_preservation": 0.20,
            "image_preservation": 0.15,
            "element_alignment": 0.15,
            "page_structure": 0.15,
        }
    else:
        weights = {
            "layout_similarity": 0.25,
            "text_preservation": 0.30,
            "image_preservation": 0.10,
            "element_alignment": 0.20,
            "page_structure": 0.15,
        }

    scores = {
        "layout_similarity": round(layout_similarity, 1),
        "text_preservation": round(text_preservation, 1),
        "image_preservation": round(image_preservation, 1),
        "element_alignment": round(element_alignment, 1),
        "page_structure": round(page_structure, 1),
    }
    overall = sum(scores[k] * weights[k] for k in scores)

    grade = (
        "A"
        if overall >= 85
        else "B"
        if overall >= 70
        else "C"
        if overall >= 55
        else "D"
    )

    return {
        "pipeline_step": "quality_score",
        "mode_used": mode_used,
        "overall": round(overall, 1),
        "grade": grade,
        # Primary explainable metrics
        "layout_similarity": scores["layout_similarity"],
        "text_preservation": scores["text_preservation"],
        "image_preservation": scores["image_preservation"],
        "element_alignment": scores["element_alignment"],
        "page_structure": scores["page_structure"],
        # Backward-compatible aliases
        "layout_fidelity": scores["layout_similarity"],
        "text_completeness": scores["text_preservation"],
        "editability": round(
            22.0 if mode_used == "visual" else _clamp(92.0 - 50.0 * spaced), 1
        ),
        "weights": weights,
        "explain": {
            "layout_similarity": layout_detail,
            "text_preservation": text_detail,
            "image_preservation": image_detail,
            "element_alignment": align_detail,
            "page_structure": structure_detail,
        },
        "thresholds": {
            "layout_retry": LAYOUT_RETRY_THRESHOLD,
            "overall_retry": OVERALL_RETRY_THRESHOLD,
        },
    }


def pick_best_candidate(
    candidates: list[dict[str, Any]],
    layout: dict[str, Any],
) -> dict[str, Any]:
    """Choose winner; prefer editable on simple docs; retry visual if layout collapses."""
    if not candidates:
        raise ValueError("No conversion candidates")

    layout_type = layout.get("layout_type", "simple")
    visual = next((c for c in candidates if c["mode"] == "visual"), None)
    editable = next((c for c in candidates if c["mode"] == "editable"), None)

    # Simple flowing docs: keep editable when reconstruction is healthy
    if layout_type == "simple" and editable:
        eq = editable["quality"]
        if (
            eq.get("layout_similarity", 0) >= 60
            and eq.get("text_preservation", 0) >= 75
            and eq.get("overall", 0) >= 65
        ):
            return editable

    def rank(c: dict[str, Any]) -> float:
        q = c["quality"]
        layout_s = q.get("layout_similarity", 0)
        text_s = q.get("text_preservation", 0)
        if layout_type in {"designed", "complex", "scanned"}:
            return 0.55 * layout_s + 0.25 * q["overall"] + 0.20 * text_s
        bonus = 3.0 if c["mode"] == "editable" and q["overall"] >= 70 else 0.0
        return q["overall"] + bonus

    best = max(candidates, key=rank)

    # Auto-retry: non-visual won but layout collapsed → switch to visual
    layout_s = best["quality"].get("layout_similarity", 0)
    if (
        visual
        and best["mode"] != "visual"
        and (
            layout_s < LAYOUT_RETRY_THRESHOLD
            or best["quality"]["overall"] < OVERALL_RETRY_THRESHOLD
        )
        and visual["quality"].get("layout_similarity", 0) > layout_s
    ):
        return {
            **visual,
            "retried_from": best["mode"],
            "retry_reason": (
                f"layout_similarity {layout_s} < {LAYOUT_RETRY_THRESHOLD} "
                f"or overall {best['quality']['overall']} < {OVERALL_RETRY_THRESHOLD}"
            ),
        }

    return best
