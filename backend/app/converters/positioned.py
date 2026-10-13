"""High-fidelity editable PDF→DOCX via absolute positioning.

Strategy (closest open-source path to ~layout-perfect editable Word):
  1. Render each PDF page as a backdrop image
  2. Cover original glyph areas on the backdrop (avoid double text)
  3. Overlay editable DrawingML text boxes at PDF text coordinates

Text remains selectable/editable in Word while visual layout stays close
to the source. True 100% is not guaranteed for exotic fonts / effects,
but this is far closer than flowing pdf2docx reconstruction on Canva PDFs.
"""

from __future__ import annotations

import secrets
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsmap, qn
from docx.shared import Emu, Pt, RGBColor
from PIL import Image, ImageDraw

_PT_TO_EMU = 12700

# DrawingML / WordprocessingML namespaces used in floating textboxes
_NS = {
    "w": nsmap["w"],
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "r": nsmap["r"],
    "xml": "http://www.w3.org/XML/1998/namespace",
}


def _emu(pt: float) -> int:
    return int(round(pt * _PT_TO_EMU))


def _rgb_from_span(span: dict) -> tuple[int, int, int]:
    color = span.get("color")
    if color is None:
        return (0, 0, 0)
    if isinstance(color, int):
        return ((color >> 16) & 255, (color >> 8) & 255, color & 255)
    return (0, 0, 0)


def _sample_fill_color(img: Image.Image, box: tuple[int, int, int, int]) -> tuple[int, int, int]:
    """Sample backdrop color from the border of a text rect."""
    x0, y0, x1, y1 = box
    x0 = max(0, x0)
    y0 = max(0, y0)
    x1 = min(img.width, x1)
    y1 = min(img.height, y1)
    if x1 <= x0 or y1 <= y0:
        return (255, 255, 255)
    samples: list[tuple[int, int, int]] = []
    step = max(1, min(x1 - x0, y1 - y0) // 8)
    for x in range(x0, x1, step):
        for y in (y0, max(y0, y1 - 1)):
            samples.append(img.getpixel((x, y)))  # type: ignore[arg-type]
    for y in range(y0, y1, step):
        for x in (x0, max(x0, x1 - 1)):
            samples.append(img.getpixel((x, y)))  # type: ignore[arg-type]
    if not samples:
        return (255, 255, 255)
    # Median-ish via average of samples (robust enough for solid Canva fills)
    r = sum(s[0] for s in samples) // len(samples)
    g = sum(s[1] for s in samples) // len(samples)
    b = sum(s[2] for s in samples) // len(samples)
    return (r, g, b)


def _zero_paragraph_spacing(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag in (qn("w:spacing"), qn("w:ind")):
            p_pr.remove(child)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)


def _add_behind_picture(paragraph, image_stream: BytesIO, width_emu: int, height_emu: int) -> None:
    """Insert a full-page picture anchored behind document text."""
    run = paragraph.add_run()
    # First add as inline to register the image relationship, then convert XML
    inline_shape = run.add_picture(image_stream, width=Emu(width_emu), height=Emu(height_emu))
    inline = inline_shape._inline
    docPr = inline.docPr
    graphic = inline.graphic
    blip = graphic.graphicData.pic.blipFill.blip
    embed = blip.get(qn("r:embed"))
    docpr_id = docPr.get("id") or "1"
    docpr_name = docPr.get("name") or "backdrop"

    # Replace inline drawing with behind-doc anchor
    anchor_xml = f"""
    <wp:anchor xmlns:wp="{_NS['wp']}" xmlns:a="{_NS['a']}" xmlns:r="{_NS['r']}"
               xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"
               distT="0" distB="0" distL="0" distR="0" simplePos="0"
               relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1"
               allowOverlap="1">
      <wp:simplePos x="0" y="0"/>
      <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
      <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
      <wp:extent cx="{width_emu}" cy="{height_emu}"/>
      <wp:effectExtent l="0" t="0" r="0" b="0"/>
      <wp:wrapNone/>
      <wp:docPr id="{docpr_id}" name="{docpr_name}"/>
      <wp:cNvGraphicFramePr/>
      <a:graphic>
        <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
          <pic:pic>
            <pic:nvPicPr>
              <pic:cNvPr id="0" name="{docpr_name}"/>
              <pic:cNvPicPr/>
            </pic:nvPicPr>
            <pic:blipFill>
              <a:blip r:embed="{embed}"/>
              <a:stretch><a:fillRect/></a:stretch>
            </pic:blipFill>
            <pic:spPr>
              <a:xfrm>
                <a:off x="0" y="0"/>
                <a:ext cx="{width_emu}" cy="{height_emu}"/>
              </a:xfrm>
              <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
            </pic:spPr>
          </pic:pic>
        </a:graphicData>
      </a:graphic>
    </wp:anchor>
    """
    drawing = run._r.find(qn("w:drawing"))
    if drawing is None:
        return
    drawing.clear()
    drawing.append(parse_xml(anchor_xml))


def _escape_xml(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _add_textbox(
    paragraph,
    *,
    text: str,
    left: int,
    top: int,
    width: int,
    height: int,
    font_size_pt: float,
    rgb: tuple[int, int, int],
    bold: bool = False,
) -> None:
    """Add an absolutely positioned editable text box."""
    width = max(width, _emu(8))
    height = max(height, _emu(8))
    shape_id = secrets.randbelow(90_000) + 1000
    r, g, b = rgb
    color_hex = f"{r:02X}{g:02X}{b:02X}"
    sz = max(10, int(round(font_size_pt * 2)))  # half-points
    b_xml = '<w:b/>' if bold else ""
    safe = _escape_xml(text)

    # Slight padding so glyphs aren't clipped
    pad = _emu(1.5)
    left = max(0, left - pad)
    top = max(0, top - pad)
    width = width + pad * 2
    height = height + pad * 2

    anchor_xml = f"""
    <w:drawing xmlns:w="{_NS['w']}" xmlns:wp="{_NS['wp']}" xmlns:a="{_NS['a']}"
               xmlns:wps="{_NS['wps']}">
      <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0"
                 relativeHeight="{shape_id}" behindDoc="0" locked="0"
                 layoutInCell="1" allowOverlap="1">
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="page">
          <wp:posOffset>{left}</wp:posOffset>
        </wp:positionH>
        <wp:positionV relativeFrom="page">
          <wp:posOffset>{top}</wp:posOffset>
        </wp:positionV>
        <wp:extent cx="{width}" cy="{height}"/>
        <wp:effectExtent l="0" t="0" r="0" b="0"/>
        <wp:wrapNone/>
        <wp:docPr id="{shape_id}" name="TextBox{shape_id}"/>
        <wp:cNvGraphicFramePr/>
        <a:graphic>
          <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
            <wps:wsp>
              <wps:cNvSpPr txBox="1"/>
              <wps:spPr>
                <a:xfrm>
                  <a:off x="0" y="0"/>
                  <a:ext cx="{width}" cy="{height}"/>
                </a:xfrm>
                <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                <a:noFill/>
                <a:ln><a:noFill/></a:ln>
              </wps:spPr>
              <wps:txbx>
                <w:txbxContent>
                  <w:p>
                    <w:pPr>
                      <w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>
                    </w:pPr>
                    <w:r>
                      <w:rPr>
                        {b_xml}
                        <w:sz w:val="{sz}"/>
                        <w:szCs w:val="{sz}"/>
                        <w:color w:val="{color_hex}"/>
                      </w:rPr>
                      <w:t xml:space="preserve">{safe}</w:t>
                    </w:r>
                  </w:p>
                </w:txbxContent>
              </wps:txbx>
              <wps:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t"/>
            </wps:wsp>
          </a:graphicData>
        </a:graphic>
      </wp:anchor>
    </w:drawing>
    """
    run = paragraph.add_run()
    run._r.append(parse_xml(anchor_xml))


def _iter_text_lines(page: fitz.Page) -> list[dict]:
    """Collect line-level text with bbox / size / color."""
    lines_out: list[dict] = []
    data = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)
    for block in data.get("blocks") or []:
        if block.get("type") != 0:
            continue
        for line in block.get("lines") or []:
            spans = line.get("spans") or []
            if not spans:
                continue
            text = "".join(s.get("text", "") for s in spans)
            if not text.strip():
                continue
            bbox = line.get("bbox") or spans[0].get("bbox")
            if not bbox:
                continue
            # Dominant span for style
            dominant = max(spans, key=lambda s: abs(s.get("size") or 0) * len(s.get("text") or ""))
            size = float(dominant.get("size") or 11)
            flags = int(dominant.get("flags") or 0)
            bold = bool(flags & 2**4) or "bold" in str(dominant.get("font", "")).lower()
            lines_out.append(
                {
                    "text": text,
                    "bbox": tuple(float(x) for x in bbox),
                    "size": size,
                    "rgb": _rgb_from_span(dominant),
                    "bold": bold,
                }
            )
    return lines_out


def pdf_to_docx_positioned(src: Path, dst: Path, *, dpi: float = 180) -> Path:
    """Convert PDF to editable DOCX with near-original layout via positioned textboxes."""
    pdf = fitz.open(src)
    document = Document()
    zoom = dpi / 72.0

    try:
        for page_index, page in enumerate(pdf):
            width_pt = float(page.rect.width)
            height_pt = float(page.rect.height)
            section = document.sections[0] if page_index == 0 else document.add_section()
            section.page_width = Emu(_emu(width_pt))
            section.page_height = Emu(_emu(height_pt))
            section.left_margin = Emu(0)
            section.right_margin = Emu(0)
            section.top_margin = Emu(0)
            section.bottom_margin = Emu(0)

            lines = _iter_text_lines(page)

            # Backdrop: render page, then cover glyph areas so editable text isn't doubled
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            draw = ImageDraw.Draw(img)
            pad_px = max(2, int(round(1.2 * zoom)))
            for line in lines:
                x0, y0, x1, y1 = line["bbox"]
                box = (
                    int(x0 * zoom) - pad_px,
                    int(y0 * zoom) - pad_px,
                    int(x1 * zoom) + pad_px,
                    int(y1 * zoom) + pad_px,
                )
                fill = _sample_fill_color(img, box)
                draw.rectangle(box, fill=fill)

            buf = BytesIO()
            img.save(buf, format="PNG")
            buf.seek(0)

            paragraph = (
                document.paragraphs[0]
                if page_index == 0 and document.paragraphs
                else document.add_paragraph()
            )
            if page_index > 0:
                paragraph = document.add_paragraph()
            _zero_paragraph_spacing(paragraph)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)

            _add_behind_picture(paragraph, buf, _emu(width_pt), _emu(height_pt))

            for line in lines:
                x0, y0, x1, y1 = line["bbox"]
                _add_textbox(
                    paragraph,
                    text=line["text"],
                    left=_emu(x0),
                    top=_emu(y0),
                    width=_emu(max(8.0, x1 - x0)),
                    height=_emu(max(line["size"] * 1.25, y1 - y0)),
                    font_size_pt=line["size"],
                    rgb=line["rgb"],
                    bold=line["bold"],
                )
    finally:
        pdf.close()

    document.save(str(dst))
    return dst
